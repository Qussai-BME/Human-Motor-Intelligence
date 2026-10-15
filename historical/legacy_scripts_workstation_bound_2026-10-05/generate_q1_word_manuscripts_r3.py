#!/usr/bin/env python3
"""Render professional DOCX manuscripts from the locked R3 Q1 Markdown manuscripts and appendices."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path("<LOCAL_WORKSPACE>/human_motor_workspace/work/human-motor-intelligence-v1_1")
OUT = ROOT / "deliverables"
PAPERS = [
    {
        "key": "P4",
        "manuscript": ROOT / "papers/P4_STRICT_ZERO_CALIBRATION_EEG_EXTERNAL_VALIDATION_V2_RECOVERY_R3_Q1.md",
        "appendix": ROOT / "papers/paper_4/APPENDIX.md",
        "output": OUT / "P4_STRICT_ZERO_CALIBRATION_EEG_EXTERNAL_VALIDATION_V2_RECOVERY_R3_Q1.docx",
        "status": "Evidence status: Reproducible external OpenBMI binary-MI validation; not a direct BCI IV-2a replication.",
        "accent": "1F4E79",
    },
    {
        "key": "P5",
        "manuscript": ROOT / "papers/P5_TRANSPARENT_EEG_EMG_FUSION_EXTERNAL_VALIDATION_V2_RECOVERY_R3_Q1.md",
        "appendix": ROOT / "papers/paper_5/APPENDIX.md",
        "output": OUT / "P5_TRANSPARENT_EEG_EMG_FUSION_EXTERNAL_VALIDATION_V2_RECOVERY_R3_Q1.docx",
        "status": "Evidence status: Reproducible amended Jeong MI-only external validation; fixed transparent fusion rules did not establish benefit over EEG-only.",
        "accent": "1F4E79",
    },
    {
        "key": "P6",
        "manuscript": ROOT / "papers/P6_INITIAL_STATE_ROBUSTNESS_AND_DECODER_QUALITY_STATUS_V2_RECOVERY_R3_Q1.md",
        "appendix": ROOT / "papers/paper_6/APPENDIX.md",
        "output": OUT / "P6_INITIAL_STATE_ROBUSTNESS_AND_DECODER_QUALITY_STATUS_V2_RECOVERY_R3_Q1.docx",
        "status": "Evidence status: R2 is a reproducible MyoSim integration-fixture systems study. R3 decoder-quality execution is blocked; authentic condition traces are absent.",
        "accent": "9C5700",
    },
]

LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
IMAGE_RE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D9E2F3") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def add_page_field(paragraph) -> None:
    paragraph.add_run("Page ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def clean_markdown(text: str) -> str:
    text = text.replace("`", "")
    text = LINK_RE.sub(lambda m: f"{m.group(1)} ({m.group(2)})", text)
    return text


def add_runs(paragraph, text: str, font_size: float = 10.2) -> None:
    text = clean_markdown(text)
    parts = BOLD_RE.split(text)
    for index, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.bold = index % 2 == 1
        run.font.name = "Cambria"
        run.font.size = Pt(font_size)


def parse_table(lines: list[str]) -> tuple[list[str], list[list[str]]]:
    rows: list[list[str]] = []
    for line in lines:
        stripped = line.strip()
        if not stripped.startswith("|"):
            continue
        cells = [cell.strip() for cell in stripped.strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            continue
        rows.append(cells)
    return rows[0], rows[1:]


def add_table(document: Document, lines: list[str], accent: str) -> None:
    headers, rows = parse_table(lines)
    if not headers:
        return
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    table_size = 6.2 if len(headers) >= 7 or len(rows) > 80 else 8.0
    for cell, label in zip(header_row.cells, headers):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, accent)
        set_cell_border(cell, "FFFFFF")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(clean_markdown(label))
        run.bold = True
        run.font.name = "Aptos Display"
        run.font.size = Pt(table_size)
        run.font.color.rgb = RGBColor(255, 255, 255)
    for row_idx, raw in enumerate(rows):
        cells = table.add_row().cells
        for col_idx, cell in enumerate(cells):
            value = raw[col_idx] if col_idx < len(raw) else ""
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_idx % 2 == 1:
                set_cell_shading(cell, "F7FAFC")
            set_cell_border(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            add_runs(p, value, table_size)
    document.add_paragraph()


def add_image(document: Document, source: Path, target: str, caption: str) -> None:
    asset = (source.parent / target).resolve()
    if not asset.exists():
        paragraph = document.add_paragraph()
        paragraph.style = "Intense Quote"
        add_runs(paragraph, f"Figure asset unavailable at rendering time: {target}", 9)
        return
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(asset), width=Inches(6.2))
    cap = document.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.style = "Caption"
    add_runs(cap, caption or "Figure", 9)


def add_markdown(document: Document, source: Path, accent: str, skip_first_h1: bool = False) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    first_h1_skipped = False
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        image = IMAGE_RE.fullmatch(stripped)
        if image:
            add_image(document, source, image.group(2), image.group(1))
            index += 1
            continue
        if stripped.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(document, table_lines, accent)
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = clean_markdown(heading.group(2))
            if level == 1 and skip_first_h1 and not first_h1_skipped:
                first_h1_skipped = True
                index += 1
                continue
            word_level = min(level, 3)
            h = document.add_heading(text, level=word_level)
            h.paragraph_format.space_before = Pt(12 if word_level == 1 else 9)
            h.paragraph_format.space_after = Pt(5)
            for run in h.runs:
                run.font.name = "Aptos Display"
                run.font.color.rgb = RGBColor.from_string(accent)
            index += 1
            continue
        if stripped.startswith(">"):
            p = document.add_paragraph(style="Intense Quote")
            add_runs(p, stripped.lstrip("> "), 10)
            index += 1
            continue
        if re.match(r"^[-*+]\s+", stripped):
            p = document.add_paragraph(style="List Bullet")
            add_runs(p, re.sub(r"^[-*+]\s+", "", stripped), 10.2)
            index += 1
            continue
        if re.match(r"^\d+\.\s+", stripped):
            p = document.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+\.\s+", "", stripped), 10.2)
            index += 1
            continue
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.14
        add_runs(p, stripped, 10.2)
        index += 1


def configure_document(document: Document, short_title: str, accent: str) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    styles = document.styles
    styles["Normal"].font.name = "Cambria"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria")
    styles["Normal"].font.size = Pt(10.2)
    for style_name in ["Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"]:
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
        style.font.color.rgb = RGBColor.from_string(accent)
    styles["Title"].font.size = Pt(24)
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 3"].font.size = Pt(11)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run(short_title + "  |  V2 Recovery R3 Q1")
    header_run.font.name = "Aptos"
    header_run.font.size = Pt(8)
    header_run.font.color.rgb = RGBColor.from_string(accent)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(footer)
    for style_name in ("Caption",):
        styles[style_name].font.name = "Cambria"
        styles[style_name].font.size = Pt(8.5)


def add_cover(document: Document, title: str, paper: dict) -> None:
    document.add_paragraph().add_run().add_break(WD_BREAK.LINE)
    title_p = document.add_paragraph(style="Title")
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(12)
    add_runs(title_p, title, 24)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(20)
    run = subtitle.add_run("Research-Grade Manuscript with Supplementary Appendix")
    run.italic = True
    run.font.name = "Cambria"
    run.font.size = Pt(12)
    info = document.add_table(rows=3, cols=2)
    info.style = "Table Grid"
    values = [
        ("Author", "Qussai Adlbi"),
        ("Version", "V2 Recovery R3 Q1-standard current-state manuscript"),
        ("Evidence boundary", paper["status"]),
    ]
    for row, (label, value) in zip(info.rows, values):
        row.cells[0].width = Inches(1.35)
        set_cell_shading(row.cells[0], paper["accent"])
        set_cell_border(row.cells[0], "FFFFFF")
        set_cell_border(row.cells[1])
        r1 = row.cells[0].paragraphs[0].add_run(label)
        r1.bold = True
        r1.font.color.rgb = RGBColor(255, 255, 255)
        r1.font.name = "Aptos Display"
        r1.font.size = Pt(9)
        p2 = row.cells[1].paragraphs[0]
        add_runs(p2, value, 9)
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(18)
    note.paragraph_format.space_after = Pt(18)
    run = note.add_run("All quantitative statements and tables are reproduced from retained programme artefacts. This DOCX is a presentation format and does not supersede the versioned Markdown evidence record.")
    run.font.name = "Cambria"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(80, 80, 80)
    document.add_page_break()


def add_landscape_section(document: Document) -> None:
    section = document.add_section(WD_SECTION_START.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)


def render(paper: dict) -> None:
    document = Document()
    document.core_properties.author = "Qussai Adlbi"
    document.core_properties.last_modified_by = "Qussai Adlbi"
    title_line = next(line[2:].strip() for line in paper["manuscript"].read_text(encoding="utf-8").splitlines() if line.startswith("# "))
    configure_document(document, paper["key"], paper["accent"])
    add_cover(document, title_line, paper)
    add_markdown(document, paper["manuscript"], paper["accent"], skip_first_h1=True)
    add_landscape_section(document)
    appendix_title = document.add_heading("Supplementary Appendix", level=1)
    for run in appendix_title.runs:
        run.font.color.rgb = RGBColor.from_string(paper["accent"])
    statement = document.add_paragraph()
    statement.style = "Intense Quote"
    add_runs(statement, "The appendix below is reproduced from the locked current-state Appendix.md. It contains retained subject-/record-level tables, statistics, validation, and reproducibility detail; unavailable prediction-level artefacts are stated rather than reconstructed.", 9.5)
    add_markdown(document, paper["appendix"], paper["accent"], skip_first_h1=True)
    paper["output"].parent.mkdir(parents=True, exist_ok=True)
    document.save(paper["output"])
    print(paper["output"])


def main() -> None:
    for paper in PAPERS:
        render(paper)


if __name__ == "__main__":
    main()
